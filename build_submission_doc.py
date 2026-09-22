import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()
    
    # Page setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # ------------------ COVER PAGE ------------------
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_inst = p_inst.add_run("THE MAHARAJA SAYAJIRAO UNIVERSITY OF BARODA\n")
    run_inst.font.name = "Arial"
    run_inst.font.size = Pt(16)
    run_inst.font.bold = True
    run_inst.font.color.rgb = RGBColor(20, 41, 77)

    run_fac = p_inst.add_run("FACULTY OF TECHNOLOGY AND ENGINEERING\nDEPARTMENT OF COMPUTER SCIENCE & ENGINEERING\n")
    run_fac.font.name = "Arial"
    run_fac.font.size = Pt(13)
    run_fac.font.bold = True
    run_fac.font.color.rgb = RGBColor(0, 119, 182)

    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.add_run("―" * 45).font.color.rgb = RGBColor(230, 126, 34)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(30)
    p_title.paragraph_format.space_after = Pt(20)
    
    r_sub = p_title.add_run("LABORATORY SUBMISSION REPORT\n")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(22)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(20, 41, 77)

    r_course = p_title.add_run("BASIC WEB PROGRAMMING (BWP)\n")
    r_course.font.name = "Arial"
    r_course.font.size = Pt(15)
    r_course.font.bold = True
    r_course.font.color.rgb = RGBColor(230, 126, 34)

    r_term = p_title.add_run("Academic Term: 2025 – 2026\nFormat: Question → Source Code → Output")
    r_term.font.name = "Arial"
    r_term.font.size = Pt(11)
    r_term.font.italic = True
    r_term.font.color.rgb = RGBColor(100, 100, 100)

    # Student Details Table
    details_table = doc.add_table(rows=4, cols=2)
    details_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    details_data = [
        ("Student Name:", "Om Kumar Khatri"),
        ("Programme / Branch:", "B.E. Computer Science & Engineering"),
        ("Subject Name:", "Basic Web Programming (BWP)"),
        ("Faculty:", "Faculty of Technology & Engineering, MSU Baroda")
    ]
    for i, (label, val) in enumerate(details_data):
        row = details_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.0)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(11)
        r0.font.color.rgb = RGBColor(20, 41, 77)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.size = Pt(11)
        
        set_cell_background(c0, "F4F6F9")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)

    doc.add_page_break()

    # ------------------ TABLE OF CONTENTS / INDEX ------------------
    p_idx = doc.add_paragraph()
    r_idx = p_idx.add_run("TABLE OF CONTENTS / LAB INDEX")
    r_idx.font.name = "Arial"
    r_idx.font.size = Pt(16)
    r_idx.font.bold = True
    r_idx.font.color.rgb = RGBColor(20, 41, 77)
    p_idx.paragraph_format.space_after = Pt(15)

    base = r"c:\Users\KHATRI OM KUMAR\Desktop\final submission\final submission"

    lab_catalog = [
        # Lab 1
        ("Lab-1", "1.1", "Design side menu using list (Computers & Accessories, Price)", "lab-1.1", ["index.html"],
         "Displays a two-column box matching the manual with side categories, products, and price ranges formatted as bulleted lists with rupee symbols."),
        ("Lab-1", "1.2", "Design first page of Faculty of Technology and Engineering using table", "lab-1.2", ["index.html"],
         "Renders the complete Faculty of Technology & Engineering portal using structured HTML <table> elements (header, nav row, news, notice board, departments table, and footer)."),
        
        # Lab 2
        ("Lab-2", "2.1", "Design a simple static website with Header, Left sidebar, and Content area Home.html", "lab-2.1", ["index.html", "Home.html", "About.html", "Contact.html"],
         "Renders the wireframe structure from the lab manual: solid blue Header, blue left navigation sidebar with links, and a right content pane dynamically showing Home.html, About.html, and Contact.html."),
        ("Lab-2", "2.2", "Create a student registration form using <form>, <input>, <textarea>, <button>, <select>, <option>", "lab-2.2", ["index.html"],
         "Displays a complete student registration form capturing Name, Gender (radios), Address (textarea), Phone, Email, Hobbies (checkboxes), City, State, Country, and College Name dropdown with submit validation."),
        ("Lab-2", "2.3", "Design MSU Home page using HTML5 and include college video in it", "lab-2.3", ["index.html"],
         "Renders the MSU Baroda homepage using semantic HTML5 tags (<header>, <nav>, <main>, <section>, <article>, <figure>, <aside>, <footer>) with an embedded college video using the HTML5 <video> tag."),

        # Lab 3
        ("Lab-3", "3.1", "Design vertical navigation bar using CSS", "lab-3.1", ["index.html"],
         "Displays an olive green vertical navigation bar with block-level menu items (HOME, NEWS, CONTACT, ABOUT) and hover feedback."),
        ("Lab-3", "3.2", "Design horizontal navigation bar using CSS", "lab-3.2", ["index.html"],
         "Displays an olive green horizontal navigation bar with floating items (HOME, NEWS, CONTACT, ABOUT) centered on the page."),
        ("Lab-3", "3.3", "Create art and craft photo album using Bootstrap", "lab-3.3", ["index.html"],
         "Renders a responsive Bootstrap 5 image card gallery showing art and craft works (Pottery, Painting, Origami, Wooden crafts, etc.) with preview buttons."),
        ("Lab-3", "3.4", "Create student marksheet using Bootstrap", "lab-3.4", ["index.html"],
         "Displays a structured academic marksheet using Bootstrap tables with subject codes, marks obtained, grades, total summary, SGPA, and result status."),
        ("Lab-3", "3.5", "Create employment application form like the given specification", "lab-3.5", ["index.html"],
         "Faithfully replicates the scanned paper application form with blue header banner, drug test notice, applicant details grid, availability schedule, and an education records table."),
        ("Lab-3", "3.6", "Design home of Flipkart with navigation bar and dropdown using Bootstrap", "lab-3.6", ["index.html"],
         "Renders a complete Flipkart home page with blue branded navbar, search bar, and active Bootstrap dropdown menus for categories and user options."),

        # Lab 4
        ("Lab-4", "4.1", "Play video in web page and play/pause, change dimension (small, normal, large) using JavaScript", "lab-4.1", ["index.html"],
         "Renders a video player with interactive buttons: Play/Pause toggles playback, and dimension buttons dynamically resize the video to 320px, 560px, or 800px."),
        ("Lab-4", "4.2", "Implement Type Writer using DHTML and JavaScript", "lab-4.2", ["index.html"],
         "Simulates a terminal typewriter effect by appending text character-by-character using setTimeout, with Start and Reset buttons."),
        ("Lab-4", "4.3", "Validate username and password using JavaScript", "lab-4.3", ["index.html"],
         "Validates credentials: username cannot start with _, @, or digits, cannot be empty; password length must be between 6 and 12 characters, and both are required."),
        ("Lab-4", "4.4", "Find first and last occurrence of vowel 'a' in word and return substring between them", "lab-4.4", ["index.html"],
         "Demonstrates with word 'ajanta': identifies first occurrence of 'a' at position 1, last occurrence at position 6, and extracts the string between them ('jant')."),
        ("Lab-4", "4.5", "Handle keyboard events (key code, vowel announcement, background change on release)", "lab-4.5", ["index.html"],
         "Captures onkeydown to display key name & keyCode, announces whenever a vowel (a, e, i, o, u) is pressed, and uses onkeyup to change the page background to blue upon release."),
        ("Lab-4", "4.6", "Design webpage demonstrating snow fall using JavaScript", "lab-4.6", ["index.html"],
         "Creates a dynamic winter snowfall animation using HTML5 Canvas and JavaScript animation loop with swaying flakes, density control, and pause/resume."),

        # Lab 5
        ("Lab-5", "5.1", "jQuery method to toggle sliding up and down a <div> element on button click", "lab-5.1", ["index.html"],
         "Uses jQuery slideToggle('slow') on button click to smoothly open and close a content panel div."),
        ("Lab-5", "5.2", "Use animate() method to move a <div> element 250 pixels to the right", "lab-5.2", ["index.html"],
         "Executes jQuery animate({ left: '250px' }) to shift a square container across an interactive stage."),
        ("Lab-5", "5.3", "Use animate() to set opacity:0.4, height:500px, width:500px for <div>", "lab-5.3", ["index.html"],
         "Demonstrates animating multiple CSS properties simultaneously: opacity reduces to 0.4 while dimensions expand to 500x500 pixels."),
        ("Lab-5", "5.4", "Use animate() to set font-size of <div> to 100 pixels with slow duration", "lab-5.4", ["index.html"],
         "Smoothly transitions the text font size from 20px to 100px using jQuery animate({ fontSize: '100px' }, 'slow')."),
        ("Lab-5", "5.5", "Create cartoon animation using jQuery", "lab-5.5", ["index.html"],
         "Animates a cartoon character face across a landscape canvas with actions like jump, roll, dance, and celebrate using chained jQuery animations."),

        # Lab 6
        ("Lab-6", "6.1", "Stock market web page using AJAX with prices from .txt file", "lab-6.1", ["index.html", "stocks.txt"],
         "Uses AJAX XMLHttpRequest to asynchronously load stocks.txt containing ticker symbols, current prices, percentage changes, and renders a live financial table."),
        ("Lab-6", "6.2", "Display cricket live score using AJAX", "lab-6.2", ["index.html", "score.json"],
         "Asynchronously retrieves match data from score.json via AJAX and dynamically displays the match summary, current batsmen, bowler stats, and recent balls."),
        ("Lab-6", "6.3", "Implement Server clock using AJAX", "lab-6.3", ["index.html", "time.txt"],
         "Regularly polls the server via AJAX every second, parsing the server timestamp headers to render an accurate digital 12-hour clock with live date display."),

        # Lab 7
        ("Lab-7", "7.1", "XML document holding cricket player collection + suitable XSD schema", "lab-7.1", ["players.xml", "players.xsd", "index.html"],
         "Defines players.xml holding player name, age, batting average, and highest score, validated against a strictly typed XML Schema Definition (players.xsd)."),
        ("Lab-7", "7.2", "XML semester marksheet + XSLT to display in XHTML table form", "lab-7.2", ["marksheet.xml", "marksheet.xsl", "index.html"],
         "Transforms marksheet.xml into an XHTML report table using marksheet.xsl with subject names, max/min marks, obtained scores, and total percentages."),
        ("Lab-7", "7.3", "XML document for book collection with fields + suitable XSD schema", "lab-7.3", ["books.xml", "books.xsd", "index.html"],
         "Defines books.xml containing book name, author, ISBN number, and quantity, validated against books.xsd with regex pattern and positive integer restrictions."),
        ("Lab-7", "7.4", "XSL file to print list of products with price < 300Rs from item.xml", "lab-7.4", ["item.xml", "item.xsl", "index.html"],
         "Applies XSLT filter <xsl:if test='price &lt; 300'> on item.xml to filter and display only affordable items in a responsive XHTML table."),
        ("Lab-7", "7.5", "XML file for book record and XSL to display in table format", "lab-7.5", ["bookrecord.xml", "bookrecord.xsl", "index.html"],
         "Binds bookrecord.xml with bookrecord.xsl using <xsl:for-each> to format book title, author, edition, price, and publisher in a clean tabular view."),

        # Node.js Lab
        ("Node.js Lab", "Node.1", "Create module for area calculation in node.js for circle, rectangle, square and use in file", "node-lab.1", ["areaModule.js", "app.js", "index.html"],
         "Executes Node.js areaModule.js exporting circle, rectangle, and square area functions. app.js imports the module, calculates values, outputs to console, and serves HTML via HTTP."),
        ("Node.js Lab", "Node.2", "Read a file from server and check whether number is Armstrong using node.js", "node-lab.2", ["app.js", "number.txt", "index.html"],
         "Node.js script reads integer from number.txt using fs.readFile, computes sum of digits raised to power of digit length, confirms Armstrong status, and outputs results."),
        ("Node.js Lab", "Node.3", "Read student.json file from server and display student info for given PRN in tabular format", "node-lab.3", ["server.js", "student.json", "index.html"],
         "Node.js HTTP server reads student.json, filters record matching query PRN, and dynamically renders student details (name, branch, semester, address, subjects) in a table.")
    ]

    # Render Index Table
    idx_table = doc.add_table(rows=1, cols=4)
    idx_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = idx_table.rows[0].cells
    hdr[0].text = "Lab #"
    hdr[1].text = "Ex #"
    hdr[2].text = "Experiment Title / Description"
    hdr[3].text = "Technology"
    
    for c in hdr:
        set_cell_background(c, "14294D")
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(255, 255, 255)

    for lab, ex, title, folder, files, desc in lab_catalog:
        row = idx_table.add_row().cells
        row[0].text = lab
        row[1].text = ex
        row[2].text = title
        
        tech = "HTML/CSS"
        if "Bootstrap" in title or "Flipkart" in title:
            tech = "Bootstrap 5"
        elif "JavaScript" in title or "DHTML" in title or "vowel" in title or "snow" in title or "keyboard" in title or "Video" in title or "validation" in title:
            tech = "JavaScript"
        elif "jQuery" in title or "cartoon" in title:
            tech = "jQuery"
        elif "AJAX" in title:
            tech = "AJAX"
        elif "XML" in title or "XSL" in title:
            tech = "XML / XSLT"
        elif "node" in title.lower() or "Node" in lab:
            tech = "Node.js"
        row[3].text = tech

        row[0].width = Inches(1.1)
        row[1].width = Inches(0.6)
        row[2].width = Inches(3.9)
        row[3].width = Inches(1.2)

        for i, c in enumerate(row):
            set_cell_margins(c, 40, 40, 60, 60)
            if i in [0, 1, 3]:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in c.paragraphs[0].runs:
                r.font.size = Pt(8.5)

    doc.add_page_break()

    # ------------------ DETAILED EXPERIMENT SECTIONS ------------------
    # Format: QUESTION -> CODE -> OUTPUT
    current_lab = ""
    for idx, (lab, ex, title, folder, files, desc) in enumerate(lab_catalog, 1):
        
        # Section Header when lab changes
        if lab != current_lab:
            current_lab = lab
            p_sec = doc.add_paragraph()
            p_sec.paragraph_format.space_before = Pt(24)
            p_sec.paragraph_format.space_after = Pt(12)
            r_sec = p_sec.add_run(f"=== {current_lab.upper()} EXERCISES ===")
            r_sec.font.name = "Arial"
            r_sec.font.size = Pt(14)
            r_sec.font.bold = True
            r_sec.font.color.rgb = RGBColor(230, 126, 34)

        # Exercise Title
        p_ex = doc.add_paragraph()
        p_ex.paragraph_format.space_before = Pt(14)
        p_ex.paragraph_format.space_after = Pt(4)
        r_ex = p_ex.add_run(f"Experiment {ex}: {title}")
        r_ex.font.name = "Arial"
        r_ex.font.size = Pt(12)
        r_ex.font.bold = True
        r_ex.font.color.rgb = RGBColor(20, 41, 77)

        # 1. QUESTION BLOCK
        p_q_lbl = doc.add_paragraph()
        p_q_lbl.paragraph_format.space_before = Pt(6)
        p_q_lbl.paragraph_format.space_after = Pt(2)
        r_qlbl = p_q_lbl.add_run("[ QUESTION STATEMENT ]")
        r_qlbl.font.name = "Arial"
        r_qlbl.font.size = Pt(9.5)
        r_qlbl.font.bold = True
        r_qlbl.font.color.rgb = RGBColor(0, 119, 182)

        q_table = doc.add_table(rows=1, cols=1)
        q_cell = q_table.rows[0].cells[0]
        q_cell.width = Inches(6.8)
        set_cell_background(q_cell, "EBF5FB")
        set_cell_margins(q_cell, 80, 80, 120, 120)
        p_q = q_cell.paragraphs[0]
        r_q = p_q.add_run(title)
        r_q.font.name = "Arial"
        r_q.font.size = Pt(9.5)
        r_q.font.bold = True
        r_q.font.color.rgb = RGBColor(20, 41, 77)

        # 2. SOURCE CODE BLOCK
        p_c_lbl = doc.add_paragraph()
        p_c_lbl.paragraph_format.space_before = Pt(10)
        p_c_lbl.paragraph_format.space_after = Pt(2)
        r_clbl = p_c_lbl.add_run("[ SOURCE CODE ]")
        r_clbl.font.name = "Arial"
        r_clbl.font.size = Pt(9.5)
        r_clbl.font.bold = True
        r_clbl.font.color.rgb = RGBColor(40, 167, 69)

        dir_path = os.path.join(base, folder)
        for fname in files:
            fpath = os.path.join(dir_path, fname)
            code_content = ""
            if os.path.exists(fpath):
                try:
                    with open(fpath, "r", encoding="utf-8", errors="replace") as f:
                        code_content = f.read()
                except Exception as e:
                    code_content = f"Error reading file: {e}"
            else:
                code_content = f"File not found: {fname}"

            # Add file banner
            p_fn = doc.add_paragraph()
            p_fn.paragraph_format.space_before = Pt(4)
            p_fn.paragraph_format.space_after = Pt(2)
            r_fn = p_fn.add_run(f"File: {folder}/{fname}")
            r_fn.font.name = "Courier New"
            r_fn.font.size = Pt(9)
            r_fn.font.bold = True
            r_fn.font.color.rgb = RGBColor(100, 100, 100)

            # Code table container (shaded box)
            code_table = doc.add_table(rows=1, cols=1)
            code_cell = code_table.rows[0].cells[0]
            code_cell.width = Inches(6.8)
            set_cell_background(code_cell, "F8F9FA")
            set_cell_margins(code_cell, 80, 80, 120, 120)
            
            p_code = code_cell.paragraphs[0]
            # Truncate if too huge for Word doc layout, but keep comprehensive
            code_display = code_content
            if len(code_display) > 7000:
                code_display = code_display[:7000] + "\n\n... [Code truncated for length. Complete source available in project file] ..."

            r_code = p_code.add_run(code_display)
            r_code.font.name = "Consolas"
            r_code.font.size = Pt(7.5)
            r_code.font.color.rgb = RGBColor(30, 30, 30)

        # 3. OUTPUT BLOCK
        p_o_lbl = doc.add_paragraph()
        p_o_lbl.paragraph_format.space_before = Pt(10)
        p_o_lbl.paragraph_format.space_after = Pt(2)
        r_olbl = p_o_lbl.add_run("[ OUTPUT & EXECUTION RESULT ]")
        r_olbl.font.name = "Arial"
        r_olbl.font.size = Pt(9.5)
        r_olbl.font.bold = True
        r_olbl.font.color.rgb = RGBColor(220, 53, 69)

        out_table = doc.add_table(rows=1, cols=1)
        out_cell = out_table.rows[0].cells[0]
        out_cell.width = Inches(6.8)
        set_cell_background(out_cell, "FDFEFE")
        set_cell_margins(out_cell, 80, 80, 120, 120)
        p_out = out_cell.paragraphs[0]
        
        r_out_title = p_out.add_run("Execution & Rendered Output Details:\n")
        r_out_title.font.bold = True
        r_out_title.font.size = Pt(9)
        r_out_title.font.color.rgb = RGBColor(20, 41, 77)

        r_out_desc = p_out.add_run(desc + "\n\n")
        r_out_desc.font.size = Pt(8.5)
        r_out_desc.font.color.rgb = RGBColor(50, 50, 50)

        r_out_path = p_out.add_run(f"Project Workspace Directory: {folder}/\nPrimary Entry File: {files[0]}")
        r_out_path.font.italic = True
        r_out_path.font.size = Pt(8)
        r_out_path.font.color.rgb = RGBColor(120, 120, 120)

        p_div2 = doc.add_paragraph()
        p_div2.paragraph_format.space_before = Pt(8)
        p_div2.paragraph_format.space_after = Pt(12)
        p_div2.add_run("―" * 60).font.color.rgb = RGBColor(210, 215, 220)

    output_path = r"C:\Users\KHATRI OM KUMAR\Desktop\Final_Lab_Submission.docx"
    doc.save(output_path)
    print(f"Master Lab Submission Document successfully created at:\n{output_path}")

if __name__ == "__main__":
    create_document()
